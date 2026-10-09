#!/usr/bin/env python3
"""A régi Excel-versenynapló nyers kinyerése JSON-ba (ADR 0048 D7, M8).

Csak olvas, és nem normalizál: a cellák értékét a típusukkal együtt
írja ki, a fejléc neve szerint. A normalizálás és a párosítás tesztelt
Dart-kódban történik (`import_legacy_races`).

Használat:

    python3 tools/legacy_race_log/xlsx_to_json.py \\
        Lola_versenynaplo_9.xlsx > legacy_races.json

Függőség: `openpyxl` (Arch: `pacman -S python-openpyxl`).
"""

import argparse
import datetime
import json
import sys

import openpyxl

FORMAT_NAME = "foretack-legacy-race-log"
FORMAT_VERSION = 1
SHEET_NAME = "Versenyek"

# A fejlécsor az, amelyben ez a két felirat áll; az első sor csak
# csoportfejléc (VERSENY, EREDMÉNY …).
DATE_HEADER = "Dátum"
NAME_HEADER = "Verseny"
HEADER_SEARCH_ROWS = 5


def cell_json(value):
    """Egy cella típusjelöléssel, vagy None, ha üres."""
    if value is None:
        return None
    # A bool az int alosztálya, ezért előbb kell vizsgálni.
    if isinstance(value, bool):
        return {"type": "text", "value": str(value)}
    if isinstance(value, (int, float)):
        return {"type": "number", "value": value}
    if isinstance(value, datetime.datetime):
        return {"type": "datetime", "value": value.isoformat()}
    if isinstance(value, datetime.date):
        return {
            "type": "datetime",
            "value": datetime.datetime.combine(
                value, datetime.time()
            ).isoformat(),
        }
    if isinstance(value, datetime.timedelta):
        return {"type": "duration", "value": value.total_seconds()}
    if isinstance(value, datetime.time):
        return {"type": "text", "value": value.isoformat()}
    return {"type": "text", "value": str(value)}


def find_header_row(sheet):
    """A fejlécsor sorszáma és a feliratai oszloponként."""
    rows = sheet.iter_rows(max_row=HEADER_SEARCH_ROWS, values_only=True)
    for number, values in enumerate(rows, start=1):
        labels = [
            value.strip() if isinstance(value, str) else None
            for value in values
        ]
        if DATE_HEADER in labels and NAME_HEADER in labels:
            named = [label for label in labels if label is not None]
            duplicates = sorted({x for x in named if named.count(x) > 1})
            if duplicates:
                raise SystemExit(f"Ismétlődő fejléc: {', '.join(duplicates)}")
            return number, labels
    raise SystemExit(
        f"Nincs „{DATE_HEADER}” és „{NAME_HEADER}” fejléc az első "
        f"{HEADER_SEARCH_ROWS} sorban."
    )


def extract_rows(sheet):
    """A fejlécek és a versenysorok (M2, M8).

    Csak az a sor verseny, amelyben a dátum és a név is ki van töltve.
    """
    header_number, labels = find_header_row(sheet)
    date_index = labels.index(DATE_HEADER)
    name_index = labels.index(NAME_HEADER)
    rows = []
    data = sheet.iter_rows(min_row=header_number + 1, values_only=True)
    for number, values in enumerate(data, start=header_number + 1):
        # Read-only módban egy sor rövidebb lehet a fejlécnél.
        def value_at(index, row=values):
            return row[index] if index < len(row) else None

        if value_at(date_index) is None or value_at(name_index) is None:
            # Az elválasztó sor és az alsó összesítő blokk itt marad ki.
            continue
        cells = {}
        for index, label in enumerate(labels):
            if label is None:
                continue
            encoded = cell_json(value_at(index))
            if encoded is not None:
                cells[label] = encoded
        rows.append({"row": number, "cells": cells})
    columns = [label for label in labels if label is not None]
    return columns, rows


def main():
    parser = argparse.ArgumentParser(
        description="Az Excel-versenynapló nyers kinyerése JSON-ba."
    )
    parser.add_argument("workbook", help="Az .xlsx fájl útvonala.")
    arguments = parser.parse_args()

    # data_only: a képlet-cellák a gyorsítótárazott értéküket adják, nem
    # a képlet szövegét (M8).
    workbook = openpyxl.load_workbook(
        arguments.workbook, data_only=True, read_only=True
    )
    if SHEET_NAME not in workbook.sheetnames:
        raise SystemExit(f"Nincs „{SHEET_NAME}” lap a munkafüzetben.")
    columns, rows = extract_rows(workbook[SHEET_NAME])
    json.dump(
        {
            "format": FORMAT_NAME,
            "version": FORMAT_VERSION,
            "sheet": SHEET_NAME,
            "columns": columns,
            "rows": rows,
        },
        sys.stdout,
        ensure_ascii=False,
        indent=1,
    )
    sys.stdout.write("\n")


if __name__ == "__main__":
    main()
